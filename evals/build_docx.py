"""Build the .docx eval cases from their text sources in evals/source/.

Two cases are Word files so the evals also exercise .docx parsing, each in a
different way real NDAs are formatted:

- 05 uses Word heading styles with no section numbers (heading-based split).
- 12 uses Word's automatic list numbering, where "1." lives outside the
  paragraph text and has to be rebuilt by the ingestion step.

In the sources, a line starting with "# " is a section heading.

    python evals/build_docx.py
"""

from pathlib import Path

import docx

HERE = Path(__file__).resolve().parent
CASES = {
    "05_disguised_non_compete": "heading_styles",
    "12_standstill": "auto_numbering",
}


def build(name: str, mode: str) -> Path:
    lines = [line.strip() for line in (HERE / "source" / f"{name}.txt").read_text(encoding="utf-8").splitlines()]
    paragraphs = [line for line in lines if line]
    d = docx.Document()
    d.add_paragraph(paragraphs[0], style="Title")
    for text in paragraphs[1:]:
        if text.startswith("# "):
            heading = text[2:]
            if mode == "heading_styles":
                d.add_heading(heading, level=1)
            else:
                d.add_paragraph(heading, style="List Number")
        else:
            d.add_paragraph(text)
    out = HERE / "ndas" / f"{name}.docx"
    d.save(out)
    return out


if __name__ == "__main__":
    for name, mode in CASES.items():
        print(build(name, mode))

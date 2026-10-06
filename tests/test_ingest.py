from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml

from nda_triage.ingest import normalize, parse_nda

SAMPLE = Path(__file__).resolve().parent.parent / "samples" / "sample_nda.txt"
W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def ids(doc):
    return [c.id for c in doc.clauses]


# ---- normalize -------------------------------------------------------------

def test_normalize_straightens_quotes_dashes_and_spaces():
    raw = "“Confidential Information” — the Party’s  data­"
    assert normalize(raw) == '"Confidential Information" - the Party\'s data'


# ---- .txt ------------------------------------------------------------------

def test_sample_txt_splits_on_numbered_sections():
    doc = parse_nda(SAMPLE)
    assert ids(doc) == [
        "preamble", "sec-1", "sec-2", "sec-2.1", "sec-2.2", "sec-3", "sec-4", "sec-5",
        "sec-6", "sec-7", "sec-8", "sec-8.1", "sec-8.2", "sec-8.3", "sec-8.4", "signature",
    ]


def test_headings_are_extracted():
    doc = parse_nda(SAMPLE)
    assert doc.clause("sec-2").heading == "Confidential Information"
    assert doc.clause("sec-4").heading == "Term"
    assert doc.clause("sec-8.3").heading == "Cooperation"


def test_lettered_subparts_stay_with_parent():
    exclusions = parse_nda(SAMPLE).clause("sec-2.2").text
    assert "(a) is or becomes publicly available" in exclusions
    assert "(d) is independently developed" in exclusions


def test_hard_wrapped_lines_are_rejoined():
    text = parse_nda(SAMPLE).clause("sec-8.3").text
    assert "neither party will solicit for employment any employee" in text


def test_offsets_trace_back_to_document_text():
    doc = parse_nda(SAMPLE)
    for c in doc.clauses:
        assert doc.text[c.start:c.end] == c.text


def test_buried_non_solicit_is_its_own_clause():
    # The trap the evals will test: a non-solicit hidden under "Cooperation".
    doc = parse_nda(SAMPLE)
    hits = [c.id for c in doc.clauses if "solicit" in c.text]
    assert hits == ["sec-8.3"]


def test_number_alone_is_not_a_section(tmp_path):
    f = tmp_path / "nda.txt"
    f.write_text("1. Term.\nThe term is two years.\n\n30 days notice is required.\n", encoding="utf-8")
    assert ids(parse_nda(f)) == ["sec-1"]


def test_unnumbered_document_falls_back_to_headings(tmp_path):
    f = tmp_path / "nda.txt"
    f.write_text("CONFIDENTIALITY\n\nKeep it secret.\n\nTerm and Termination\n\nTwo years.\n", encoding="utf-8")
    doc = parse_nda(f)
    assert [c.heading for c in doc.clauses] == ["CONFIDENTIALITY", "Term and Termination"]


def test_unsupported_type_is_rejected(tmp_path):
    f = tmp_path / "nda.pdf"
    f.write_bytes(b"%PDF")
    with pytest.raises(ValueError, match="PDF is planned"):
        parse_nda(f)


# ---- .docx -----------------------------------------------------------------

def _add_numbering(document) -> str:
    """Register a list whose levels render as '1.', '1.1', '(a)', like an NDA."""
    numbering = document.part.numbering_part.element
    numbering.append(parse_xml(
        f'<w:abstractNum {W_NS} w:abstractNumId="900">'
        '<w:lvl w:ilvl="0"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1."/></w:lvl>'
        '<w:lvl w:ilvl="1"><w:start w:val="1"/><w:numFmt w:val="decimal"/><w:lvlText w:val="%1.%2"/></w:lvl>'
        '<w:lvl w:ilvl="2"><w:start w:val="1"/><w:numFmt w:val="lowerLetter"/><w:lvlText w:val="(%3)"/></w:lvl>'
        '</w:abstractNum>'
    ))
    numbering.append(parse_xml(f'<w:num {W_NS} w:numId="900"><w:abstractNumId w:val="900"/></w:num>'))
    return "900"


def _numbered(document, text, level, num_id):
    p = document.add_paragraph(text)
    p._p.get_or_add_pPr().append(parse_xml(
        f'<w:numPr {W_NS}><w:ilvl w:val="{level}"/><w:numId w:val="{num_id}"/></w:numPr>'
    ))


def test_docx_auto_numbering_is_rebuilt(tmp_path):
    d = docx.Document()
    num_id = _add_numbering(d)
    d.add_paragraph("This Agreement is made between Acme and Birch.")
    _numbered(d, "Definitions.", 0, num_id)
    _numbered(d, "Confidential Information. Means all non-public information.", 1, num_id)
    _numbered(d, "Exclusions. Does not include:", 1, num_id)
    _numbered(d, "public information;", 2, num_id)
    _numbered(d, "independently developed information.", 2, num_id)
    _numbered(d, "Term. Two years.", 0, num_id)
    path = tmp_path / "nda.docx"
    d.save(path)

    doc = parse_nda(path)
    assert ids(doc) == ["preamble", "sec-1", "sec-1.1", "sec-1.2", "sec-2"]
    assert doc.clause("sec-1.1").heading == "Confidential Information"
    assert doc.clause("sec-1.2").text == "1.2 Exclusions. Does not include:\n(a) public information;\n(b) independently developed information."


def test_docx_heading_styles_without_numbering(tmp_path):
    d = docx.Document()
    d.add_heading("Confidentiality obligations of the receiving party", level=1)  # too long for the text heuristic
    d.add_paragraph("Keep it secret.")
    d.add_heading("Term", level=1)
    d.add_paragraph("Two years.")
    path = tmp_path / "nda.docx"
    d.save(path)

    doc = parse_nda(path)
    assert [c.heading for c in doc.clauses] == ["Confidentiality obligations of the receiving party", "Term"]


def test_docx_tables_are_read(tmp_path):
    d = docx.Document()
    d.add_paragraph("1. Parties.")
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Acme Analytics, Inc."
    table.cell(0, 1).text = "Birch Robotics LLC"
    path = tmp_path / "nda.docx"
    d.save(path)

    assert "Acme Analytics, Inc. | Birch Robotics LLC" in parse_nda(path).clause("sec-1").text
